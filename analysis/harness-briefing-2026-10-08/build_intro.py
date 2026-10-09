#!/usr/bin/env python3
"""Short external introduction deck (16:9, 15 slides) built around the four topics the user asked for:
running approach, tools, OctoLoop scheduling, evolution. Plain wording; no principle list, few incidents.

    python3 build_intro.py      # needs python-pptx; writes $OUT/harness-intro.pptx
                                # (default $WORKSPACES/westlake-inputs/reports/2026-10-08-harness-briefing)
Images: intro-lit / intro-white / intro-err (U14 signed t20 shots: Wikipedia lit, OrganicMaps alive-white,
SuperTuxKart error dialog), copied next to this script."""
import os
import sys
from pathlib import Path
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
from pptx.oxml.ns import qn

HERE = Path(__file__).resolve().parent
sys.path.insert(0, str(HERE.parents[1] / 'scripts' / 'lab'))
import lab_paths  # noqa: E402
OUT = Path(os.environ.get('OUT') or f'{lab_paths.workspaces()}/westlake-inputs/reports/2026-10-08-harness-briefing')
OUT.mkdir(parents=True, exist_ok=True)

FONT = 'PingFang SC'
BLUE, DARK, GREY = RGBColor(0x0a, 0x6c, 0xff), RGBColor(0x1d, 0x1d, 0x1f), RGBColor(0x6e, 0x6e, 0x73)
LIGHT, WHITE = RGBColor(0xf3, 0xf6, 0xfb), RGBColor(0xff, 0xff, 0xff)
GREEN, RED, ORANGE, PURPLE = RGBColor(0x1f, 0x9d, 0x55), RGBColor(0xd7, 0x26, 0x3d), RGBColor(0xf0, 0x8c, 0x00), RGBColor(0x6e, 0x4b, 0xd8)
PALE = {'blue': RGBColor(0xe3, 0xee, 0xff), 'green': RGBColor(0xe8, 0xf6, 0xec), 'grey': RGBColor(0xee, 0xee, 0xf0),
        'orange': RGBColor(0xff, 0xf4, 0xe0), 'purple': RGBColor(0xf0, 0xea, 0xff), 'red': RGBColor(0xff, 0xec, 0xec)}
prs = Presentation(); prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
BL = prs.slide_layouts[6]


def _font(run, size, bold, color):
    f = run.font; f.size = Pt(size); f.bold = bold; f.color.rgb = color; f.name = FONT
    rPr = run._r.get_or_add_rPr(); ea = rPr.find(qn('a:ea'))
    if ea is None:
        ea = rPr.makeelement(qn('a:ea'), {}); rPr.append(ea)
    ea.set('typeface', FONT)


def text(s, x, y, w, h, lines, size=16, bold=False, color=DARK, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP, spacing=6):
    tb = s.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h)); tf = tb.text_frame; tf.word_wrap = True
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = Emu(0); tf.vertical_anchor = anchor
    for i, ln in enumerate(lines.split('\n') if isinstance(lines, str) else lines):
        t, o = (ln, {}) if isinstance(ln, str) else ln
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph(); p.alignment = align; p.space_after = Pt(spacing)
        r = p.add_run(); r.text = ('• ' + t) if o.get('bullet') else t
        _font(r, o.get('size', size), o.get('bold', bold), o.get('color', color))
    return tb


def box(s, x, y, w, h, fill=LIGHT, line=None, shape=MSO_SHAPE.ROUNDED_RECTANGLE):
    b = s.shapes.add_shape(shape, Inches(x), Inches(y), Inches(w), Inches(h))
    b.fill.solid(); b.fill.fore_color.rgb = fill; b.shadow.inherit = False
    if line is None:
        b.line.fill.background()
    else:
        b.line.color.rgb = line; b.line.width = Pt(1.5)
    if shape == MSO_SHAPE.ROUNDED_RECTANGLE:
        b.adjustments[0] = 0.08
    return b


def card(s, x, y, w, h, title, body=None, fill=LIGHT, tsize=17, bsize=13, tcolor=DARK, align=PP_ALIGN.LEFT, line=None):
    box(s, x, y, w, h, fill, line)
    text(s, x + 0.2, y + 0.15, w - 0.4, 0.5, title, tsize, True, tcolor, align=align)
    if body:
        text(s, x + 0.2, y + 0.2 + tsize / 72 * 1.7, w - 0.4, h - 0.6, body, bsize, False, DARK, align=align, spacing=4)


def arrow(s, x1, y1, x2, y2, color=GREY, width=2.0):
    c = s.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(x1), Inches(y1), Inches(x2), Inches(y2))
    c.line.color.rgb = color; c.line.width = Pt(width)
    ln = c.line._get_or_add_ln(); ln.append(ln.makeelement(qn('a:tailEnd'), {'type': 'triangle', 'w': 'med', 'len': 'med'}))


def slide(title, part=None, sub=None):
    s = prs.slides.add_slide(BL)
    bar = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, prs.slide_width, Inches(0.08)); bar.fill.solid(); bar.fill.fore_color.rgb = BLUE; bar.line.fill.background()
    if part:
        text(s, 0.5, 0.25, 8, 0.35, part, 12, True, BLUE)
    text(s, 0.5, 0.5 if part else 0.35, 12.3, 0.7, title, 28, True)
    if sub:
        text(s, 0.5, 1.15 if part else 1.0, 12.3, 0.45, sub, 14, False, GREY)
    text(s, 12.4, 7.05, 0.6, 0.3, str(len(prs.slides)), 10, False, GREY, align=PP_ALIGN.RIGHT)
    return s


def table(s, x, y, w, rows, colw, size=13, row_h=0.5):
    t = s.shapes.add_table(len(rows), len(rows[0]), Inches(x), Inches(y), Inches(w), Inches(row_h * len(rows))).table
    for j, cw in enumerate(colw):
        t.columns[j].width = Inches(cw)
    for i, row in enumerate(rows):
        for j, v in enumerate(row):
            c = t.cell(i, j); c.text = ''; c.margin_left = c.margin_right = Inches(0.1); c.margin_top = c.margin_bottom = Inches(0.05)
            r = c.text_frame.paragraphs[0].add_run(); r.text = v
            _font(r, size, i == 0, WHITE if i == 0 else DARK)
            c.fill.solid(); c.fill.fore_color.rgb = BLUE if i == 0 else (WHITE if i % 2 else LIGHT)


P1, P2, P3, P4 = '一 · 运行思路', '二 · 工具', '三 · OctoLoop 调度', '四 · 进化'

# 1 cover ---------------------------------------------------------------------------------------------------------
s = prs.slides.add_slide(BL)
box(s, 0, 0, 13.333, 7.5, RGBColor(0x0b, 0x1f, 0x3a), shape=MSO_SHAPE.RECTANGLE)
text(s, 0.9, 2.1, 11.5, 1.0, 'Westlake harness', 44, True, WHITE)
text(s, 0.9, 3.1, 11.5, 0.6, '用多个 AI 车道,把 Android 应用原包跑上 OpenHarmony', 24, False, RGBColor(0xc8, 0xd6, 0xf0))
text(s, 0.9, 4.4, 11.5, 0.5, '运行思路 · 工具 · OctoLoop 调度 · 进化', 18, False, RGBColor(0x9f, 0xb3, 0xd6))

# 2 where we are ----------------------------------------------------------------------------------------------------
s = slide('现在做到哪')
card(s, 0.5, 1.4, 4.0, 2.6, '128 个应用能上屏', '未修改的 Android APK 在 OH 6.1 开发板上安装、启动,约第 20 秒的截图里是应用自己的界面。', PALE['blue'], 20, 15)
card(s, 4.67, 1.4, 4.0, 2.6, '多数只到「能上屏」', '登录、网络、相机、后台等功能还没有逐项验收。上屏只是第一道门槛。', PALE['orange'], 20, 15)
card(s, 8.83, 1.4, 4.0, 2.6, '今日头条走得最远', '早期路线上已经可用:能刷首页、能点开文章;预编译加 JIT 把点开文章从 14–24 秒降到 0.83 秒。约 15–20% 冷启动会崩,靠看门狗自动拉起。当前原包路线上还在攻一堵原生墙。', PALE['green'], 20, 13)
text(s, 0.5, 4.5, 12.3, 0.5, '今天讲四件事', 18, True)
for i, (t, b) in enumerate([('运行思路', '横向推进,先修高价值的 API'), ('工具', '静态分析、环境检查、怎么确认'), ('OctoLoop 调度', '一个外环带多个内环'), ('进化', '一轮之后改流程、改工具')]):
    x = 0.5 + i * 3.13
    box(s, x, 5.1, 2.95, 1.5, LIGHT)
    text(s, x + 0.2, 5.2, 0.6, 0.6, str(i + 1), 30, True, BLUE)
    text(s, x + 0.8, 5.3, 2.0, 0.5, t, 17, True)
    text(s, x + 0.8, 5.8, 2.05, 0.7, b, 12, False, GREY)

# 3 horizontal vs single --------------------------------------------------------------------------------------------
s = slide('横向推进,不单兵推进', P1, '同一堵墙往往挡住很多应用;一次修在公共 API 上,比一个应用一个应用地硬过要稳')
box(s, 0.5, 1.8, 5.6, 4.6, PALE['grey'])
text(s, 0.75, 1.95, 5.2, 0.5, '单兵推进(不这么做)', 19, True, GREY)
for i, t in enumerate(['盯住一个应用', '哪里报错改哪里', '为了过而过,容易乱改', '改动只对这一个应用有效,还可能弄坏别的']):
    y = 2.65 + i * 0.92
    box(s, 0.85, y, 4.9, 0.7, WHITE)
    text(s, 1.0, y + 0.17, 4.6, 0.4, t, 15, False, GREY)
    if i < 3:
        arrow(s, 3.3, y + 0.7, 3.3, y + 0.92, GREY)
box(s, 6.6, 1.8, 6.2, 4.6, PALE['blue'])
text(s, 6.85, 1.95, 5.8, 0.5, '横向推进(现在的做法)', 19, True, BLUE)
steps = ['全量跑一遍所有应用,记录每个应用的首个致命错误', '按错误归类,挑挡住应用最多的那个 API', '在公共层修一次(API shim)', '在多个应用上验证同一个修复', '验证过就冻结,以后不许悄悄改']
for i, t in enumerate(steps):
    y = 2.55 + i * 0.75
    c = s.shapes.add_shape(MSO_SHAPE.OVAL, Inches(6.95), Inches(y + 0.07), Inches(0.42), Inches(0.42)); c.fill.solid(); c.fill.fore_color.rgb = BLUE; c.line.fill.background()
    text(s, 6.95, y + 0.13, 0.42, 0.3, str(i + 1), 13, True, WHITE, align=PP_ALIGN.CENTER)
    text(s, 7.55, y + 0.1, 5.1, 0.5, t, 15)
text(s, 0.5, 6.6, 12.3, 0.6, '例:安装器给 Android 应用加一个「允许后台启动」的授权,只改一处,Tusky、K-9 Mail、Thunderbird、文件管理器 4 个应用都从白屏变成了自己的界面,随后登记冻结。', 13, False, GREY)

# 4 one round -------------------------------------------------------------------------------------------------------
s = slide('一轮怎么推进', P1, '三块板平时保持同一个版本(统一态);每一轮只出一个候选版本,把这一轮所有修复合在一起')
flow = [('统一态', '三块板同一版本,\n逐件哈希一致'), ('全量观测', '每个应用截图 + 日志'), ('归类选墙', '按首个致命错误归类,\n先修挡住最多的'), ('合成候选', '本轮修复按层合成\n一个版本'), ('观察窗', '三板分片验证,\n窗末一律回滚'), ('晋级', '无回退才成为\n新统一态')]
for i, (t, b) in enumerate(flow):
    x = 0.45 + i * 2.12
    card(s, x, 1.9, 1.85, 1.9, t, b, PALE['blue'] if i not in (4, 5) else PALE['green'], 17, 12, align=PP_ALIGN.CENTER)
    if i < 5:
        arrow(s, x + 1.85, 2.85, x + 2.12, 2.85, BLUE)
arrow(s, 0.45 + 5 * 2.12 + 0.92, 3.8, 0.45 + 5 * 2.12 + 0.92, 4.25, GREEN)
arrow(s, 0.45 + 5 * 2.12 + 0.92, 4.25, 1.37, 4.25, GREEN)
arrow(s, 1.37, 4.25, 1.37, 3.8, GREEN)
text(s, 4.2, 4.32, 5, 0.4, '下一轮从新的统一态开始', 12, False, GREEN, align=PP_ALIGN.CENTER)
card(s, 0.5, 4.95, 6.0, 1.9, '什么时候冻结', '同一个公共 API 的修复,至少 2 个不同应用过了这堵墙,其中至少 1 个截图确认上屏,就登记冻结。冻结后每次构建、部署前自动核对。', PALE['purple'], 17, 13)
card(s, 6.83, 4.95, 6.0, 1.9, '什么时候不晋级', '观察窗里出现确认的回退,或者新版本没有实际收益(只是和旧版打平),就不晋级,问题对半拆开再查。', PALE['orange'], 17, 13)

# 5 tools overview --------------------------------------------------------------------------------------------------
s = slide('工具分三类', P2, '仓库里自己写的工具 137 个,另有 122 个测试和回放;下面按用途分')
cols = [('静态分析', '不碰板子', '读 APK、二进制、历史日志,回答:\n里面有什么?库开进了哪里?会挡住哪些应用?修复改动了什么?', PALE['blue'], BLUE),
        ('环境检查', '核对真实状态', '板子、部署件、构建环境、控制机,回答:\n现在跑的到底是哪一版?修复真的生效了吗?环境和上次一样吗?', PALE['green'], GREEN),
        ('执行与调度', '动手的部分', '上板、事务式部署和回滚、派单、侦听、代为入库。\n所有写操作都走这里,并且都带回滚。', PALE['orange'], ORANGE)]
for i, (t, sub, b, fill, col) in enumerate(cols):
    x = 0.5 + i * 4.18
    box(s, x, 1.8, 3.95, 3.3, fill)
    text(s, x + 0.25, 1.95, 3.5, 0.5, t, 22, True, col)
    text(s, x + 0.25, 2.5, 3.5, 0.4, sub, 13, False, GREY)
    text(s, x + 0.25, 3.0, 3.5, 2.0, b, 14, spacing=6)
box(s, 0.5, 5.35, 12.33, 1.5, LIGHT)
text(s, 0.75, 5.5, 11.9, 0.4, '工具怎么来的', 17, True)
text(s, 0.75, 5.95, 11.9, 0.9, ['每次上板出了问题,都追问一句:哪个检查本可以在上板前拦住它?答案就做成工具或门禁,接进流水线。',
                                 '另外两部分:收编原有的实验室脚本;为验收契约写的测试(让 agent 的交付能被机器检查)。'], 14, spacing=4)

# 6 static analysis -------------------------------------------------------------------------------------------------
s = slide('静态分析工具', P2, '凡是离线能回答的问题,不占用板子')
table(s, 0.5, 1.7, 12.33, [
    ('工具', '回答什么问题'),
    ('open_site_scan', '宿主和子进程里每个库从哪里加载、进了哪个 namespace、会不会被加载两份'),
    ('static_pipeline / aggregate_gaps', '全部 APK 缺哪些系统服务、JNI、库;哪个缺口挡住的应用最多'),
    ('first_fatal / cluster_first_cause', '每个应用的首个致命错误是什么;按挡住的应用数排序,决定先修哪堵墙'),
    ('predict_vs_actual', '全量前写下每个应用的预测,全量后对账;没预测到的墙补进扫描器'),
    ('abi_diff / deployed_class_diff / deployed_strings_diff', '新产物和板上正在跑的那份比:导出函数、调用签名、类、字符串有没有被删或改'),
    ('null_producer', '空指针来自框架接口的返回值,还是应用自己的代码'),
    ('REA + Hopper(新接入)', '反编译 native 库,辅助判断崩溃点;结论要回到原始指令再核对一遍'),
], [4.4, 7.93], size=13, row_h=0.6)
text(s, 0.5, 6.7, 12.3, 0.5, '静态计数只是上界:比如某个符号 50 个应用引用,实际因它失败的只有 2 个。排序最终以实测的首个致命错误为准。', 13, False, GREY)

# 7 environment checks ----------------------------------------------------------------------------------------------
s = slide('环境检查工具:种类多,目的是一个', P2, '确认「现在跑的就是我们以为的那一版」,环境一变,结论就不能沿用')
groups = [('板子状态', ['设备在线、锁是否被占用', '每条远端命令都带结束标记并回读', '运行时指纹:板上每个件的哈希', '板上比出厂多了什么、安装链身份']),
          ('部署件身份', ['部署时逐件核对哈希', '激活冒烟:修复在板上真的生效', '冻结检查:冻结件没有被改动', '实际加载的是哪一份库/类']),
          ('构建与工具链', ['工具链和补丁序列的哈希', '编译器与 boot image 可复现', '镜像与 jar 一致、编译入口对照', 'native 预部署门、JNI 覆盖门']),
          ('控制机与路径', ['Mac 和虚机两面跑同一套门,结果不一致就停', '脚本不写死个人路径', '外部构建机不通就停在加锁前,不绕门禁'])]
for i, (t, items) in enumerate(groups):
    x, y = 0.5 + (i % 2) * 6.22, 1.75 + (i // 2) * 2.6
    box(s, x, y, 6.0, 2.4, PALE['green'])
    text(s, x + 0.25, y + 0.15, 5.5, 0.4, t, 18, True, GREEN)
    text(s, x + 0.25, y + 0.65, 5.5, 1.7, [(it, {'bullet': True}) for it in items], 14, spacing=4)

# 8 how we confirm --------------------------------------------------------------------------------------------------
s = slide('怎么确认真的点亮', P2, '进程活着、脚本打印 PASS 都不算;以截图为准,再用其他证据补充')
for i, (k, cap, col) in enumerate([('intro-lit', '算:应用自己的界面', GREEN), ('intro-white', '不算:进程活着,但白屏', RED), ('intro-err', '单列:有界面,但带报错框', ORANGE)]):
    x = 0.5 + i * 2.3
    pic = s.shapes.add_picture(str(HERE / f'{k}.jpg'), Inches(x), Inches(1.75), height=Inches(3.4)); pic.line.color.rgb = RGBColor(0x9a, 0x9a, 0x9e); pic.line.width = Pt(1.5)
    fr = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, pic.left, pic.top, pic.width, pic.height); fr.fill.background()
    fr.line.color.rgb = RGBColor(0x9a, 0x9a, 0x9e); fr.line.width = Pt(1.5); fr.shadow.inherit = False
    text(s, x, 5.25, 2.15, 0.6, cap, 12, True, col)
text(s, 7.6, 1.75, 5.2, 5.2, [
    ('截图', {'bold': True, 'size': 16}),
    ('启动后约第 20 秒截图,外环逐张看图签认。这是最终依据。', {}),
    ('计数从记录里数', {'bold': True, 'size': 16}),
    ('截了几张、进程是否存活,由工具从运行记录里数出来,汇报只能贴工具输出。', {}),
    ('和参考配对', {'bold': True, 'size': 16}),
    ('新版截图和上一版同一应用的截图配对比较;不稳定的应用重复跑多次,看几次里亮几次。', {}),
    ('日志与自检', {'bold': True, 'size': 16}),
    ('首个致命错误消失或后移,说明修复生效;功能级的回退(比如加密出错)截图看不出,靠进程自检拦住。', {}),
], 13, spacing=5)

# 9 freeze / knowledge / reproduce ----------------------------------------------------------------------------------
s = slide('冻结、积累知识、可复现', P2)
cols = [('冻结', PALE['purple'], PURPLE, ['登记:API、实现文件的哈希、验证过的应用和截图', '每次构建、部署前自动核对,不符就拒', '只允许三种改动:证明原实现有缺陷、平台变化必须重编、加功能且原行为不变', '削弱或删除冻结项,只有项目负责人能批', '目前 8 项']),
        ('积累知识', PALE['blue'], BLUE, ['确认的事实写进知识摘要,操作方法写进操作手册', '每条都带证据路径', '写了「根因」却没有对照实验的,工具直接拒收', '只放在某个 agent 自己记忆里的,不算']),
        ('可复现', PALE['green'], GREEN, ['每个结论一个证据目录:说明、结果、复现脚本', '上板的产物连同源码、补丁、工具链哈希、构建日志一起保存', '关键检查能在留存数据上重放', '构建机换了,修复也不会丢'])]
for i, (t, fill, col, items) in enumerate(cols):
    x = 0.5 + i * 4.18
    box(s, x, 1.4, 3.95, 4.4, fill)
    text(s, x + 0.25, 1.55, 3.5, 0.5, t, 20, True, col)
    text(s, x + 0.25, 2.2, 3.5, 3.5, [(it, {'bullet': True}) for it in items], 15, spacing=9)

box(s, 0.5, 6.05, 12.33, 0.85, LIGHT)
text(s, 0.75, 6.25, 11.9, 0.5, '三件事互相支撑:冻结保证修好的不被改坏,知识让下一个会话不用重新摸索,可复现让结论经得起别人重做。', 15, False, DARK)

# 10 OctoLoop diagram -----------------------------------------------------------------------------------------------
s = slide('一个外环,带多个内环', P3, '外环不写业务代码,负责派单、复验、看图签认、入库;内环各有专长,各自在独立的工作目录里干活')
card(s, 4.65, 1.6, 4.0, 0.9, '项目负责人', '方向、优先级、保留的审批', PALE['orange'], 15, 11, align=PP_ALIGN.CENTER)
card(s, 0.5, 3.0, 3.4, 2.0, '外环', '派单 · 复验 · 看图签认\n代为入库 · 写知识', PALE['blue'], 19, 13, BLUE, align=PP_ALIGN.CENTER)
card(s, 4.65, 3.0, 4.0, 2.0, '黑板', '所有派单、进度、交件、板锁\n都写在这里,按固定格式\n(唯一事实源)', WHITE, 19, 13, BLUE, align=PP_ALIGN.CENTER, line=BLUE)
lanes = [('板上执行', 'codex'), ('静态分析与门禁', 'codex'), ('native 与反编译', 'Claude'), ('WebView / ArkWeb', 'Claude'), ('机械扫描', 'octoscode')]
for i, (t, who) in enumerate(lanes):
    y = 1.6 + i * 0.78
    box(s, 9.4, y, 3.45, 0.66, PALE['green'])
    text(s, 9.6, y + 0.1, 2.2, 0.45, t, 13, True)
    text(s, 11.75, y + 0.13, 1.0, 0.4, who, 11, False, GREY, align=PP_ALIGN.RIGHT)
text(s, 9.4, 5.5, 2.5, 0.4, '内环 5 条车道', 12, True, GREEN)
card(s, 0.5, 5.75, 5.2, 1.2, '仓库', '代码、工具、知识、证据、冻结登记', PALE['purple'], 16, 12, align=PP_ALIGN.CENTER)
card(s, 7.6, 6.0, 5.25, 1.0, '3 块 OH 开发板', '加锁使用 · 用完回滚', PALE['red'], 15, 11, align=PP_ALIGN.CENTER)
arrow(s, 6.65, 2.5, 6.65, 3.0, ORANGE)
arrow(s, 3.9, 3.7, 4.65, 3.7, BLUE); arrow(s, 4.65, 4.3, 3.9, 4.3, GREEN)
arrow(s, 8.65, 3.7, 9.4, 3.2, BLUE); arrow(s, 9.4, 4.5, 8.65, 4.3, GREEN)
arrow(s, 2.2, 5.0, 2.2, 5.75, PURPLE); arrow(s, 12.3, 5.4, 12.3, 6.0, RED)
text(s, 3.9, 5.1, 5.5, 0.6, ['蓝:派单   绿:进度与交件', '紫:入库   红:上板(加锁)'], 11, False, GREY, align=PP_ALIGN.CENTER, spacing=2)

# 11 scheduling -----------------------------------------------------------------------------------------------------
s = slide('怎么合理调度', P3)
steps = [('拆单', '每个任务一个编号,写清目标、交付物、验收标准、时间预算;验收尽量写成机器能检查的条件。'),
         ('按专长派', '板上事务、静态分析、native 逆向、WebView、机械扫描各有车道;某类任务某个模型会拒绝,就派给另一个。'),
         ('先离线,后上板', '静态能回答的问题不占板子;上板前先写预测,上板后对账。'),
         ('板子是稀缺资源', '同一块板同一时刻只有一个任务;先在黑板宣布,再加锁;单项不超过 30 分钟;用完回滚、放锁。'),
         ('交件要复验', '车道交件清单,外环独立复验,再经门禁和测试代为入库;计数类结论只认工具输出。')]
for i, (t, b) in enumerate(steps):
    y = 1.35 + i * 1.12
    c = s.shapes.add_shape(MSO_SHAPE.OVAL, Inches(0.5), Inches(y + 0.1), Inches(0.55), Inches(0.55)); c.fill.solid(); c.fill.fore_color.rgb = BLUE; c.line.fill.background()
    text(s, 0.5, y + 0.2, 0.55, 0.4, str(i + 1), 16, True, WHITE, align=PP_ALIGN.CENTER)
    text(s, 1.3, y + 0.12, 3.0, 0.5, t, 18, True)
    text(s, 4.3, y + 0.12, 8.5, 0.9, b, 15)

# 12 long running ---------------------------------------------------------------------------------------------------
s = slide('怎么做到长期运行', P3, '人不在的时候也能连续推进;出了问题停在安全的地方,而不是硬闯')
loop = [('派单', 3.0, 1.75), ('车道执行', 6.2, 1.75), ('交件 / 报阻塞', 6.2, 3.55), ('外环复验、入库', 3.0, 3.55)]
for t, x, y in loop:
    card(s, x, y, 2.6, 0.95, t, None, PALE['blue'], 16, align=PP_ALIGN.CENTER)
arrow(s, 5.6, 2.22, 6.2, 2.22, BLUE); arrow(s, 7.5, 2.7, 7.5, 3.55, BLUE); arrow(s, 6.2, 4.02, 5.6, 4.02, BLUE); arrow(s, 4.3, 3.55, 4.3, 2.7, BLUE)
text(s, 4.2, 2.95, 3.4, 0.5, '黑板记录每一步', 13, True, BLUE, align=PP_ALIGN.CENTER)
card(s, 9.4, 1.75, 3.45, 2.75, '后台哨兵', '新交件哨:黑板出现交件就唤醒外环\n车道状态哨:车道停摆、断连、卡住就报警\n每小时巡检兜底', PALE['green'], 16, 13)
box(s, 0.5, 5.0, 12.33, 1.95, LIGHT)
text(s, 0.75, 5.12, 11.9, 0.4, '出问题时怎么办', 17, True)
text(s, 0.75, 5.6, 5.8, 1.3, [('车道额度用完或卡住 → 转派给别的车道', {'bullet': True}), ('板子掉线 → 按事务记录回滚、恢复,核对后才放锁', {'bullet': True})], 14, spacing=6)
text(s, 6.8, 5.6, 5.9, 1.3, [('外部构建机不通 → 停在加锁前等待,不绕过检查', {'bullet': True}), ('会话断了 → 状态都在黑板和仓库里,新会话能接着做', {'bullet': True})], 14, spacing=6)

# 13 evolution cycle ------------------------------------------------------------------------------------------------
s = slide('一轮之后怎么进化', P4, '先改流程、改工具,再积累知识、推进代码;最后由下一轮检验这一轮的问题是不是真的解决了')
stages = [('改流程', '这一轮哪一步反复出错,\n就改成规则', BLUE, PALE['blue'], True),
          ('改工具', '把规则做成工具或门禁,\n让它自动拦住', BLUE, PALE['blue'], True),
          ('积累知识', '确认的事实、操作方法、\n验证过的 API 冻结', PURPLE, PALE['purple'], False),
          ('推进代码', '本轮修复合成\n下一版候选', ORANGE, PALE['orange'], False),
          ('下一轮验证', '对照本轮的预测和问题清单,\n看是不是真解决了', GREEN, PALE['green'], False)]
for i, (t, b, col, fill, key) in enumerate(stages):
    x = 0.4 + i * 2.55
    box(s, x, 1.95, 2.25, 2.7, fill, col if key else None)
    text(s, x + 0.15, 2.1, 1.95, 0.5, str(i + 1) + ' ' + t, 19, True, col, align=PP_ALIGN.CENTER)
    text(s, x + 0.15, 2.8, 1.95, 1.7, b, 13, align=PP_ALIGN.CENTER)
    if key:
        text(s, x + 0.15, 4.15, 1.95, 0.35, '重点', 12, True, col, align=PP_ALIGN.CENTER)
    if i < 4:
        arrow(s, x + 2.25, 3.3, x + 2.55, 3.3, GREY)
arrow(s, 0.4 + 4 * 2.55 + 1.12, 4.65, 0.4 + 4 * 2.55 + 1.12, 5.05, GREEN)
arrow(s, 0.4 + 4 * 2.55 + 1.12, 5.05, 1.52, 5.05, GREEN)
arrow(s, 1.52, 5.05, 1.52, 4.65, GREEN)
text(s, 4.0, 5.12, 5.5, 0.4, '没解决的问题进入下一轮', 12, False, GREEN, align=PP_ALIGN.CENTER)
text(s, 0.5, 5.75, 12.3, 1.2, [('为什么流程和工具排在前面:知识会过时,代码会被改掉,但一条好的规则和一个能自动拦截的工具,会在之后每一轮里持续起作用。', {}),
                               ('例:汇报里的截图数和存活数曾经按计划去写而不是数出来;现在规定只能贴工具从记录里数出的结果,这一类错误之后改由工具把关,不再靠人自觉。', {'color': GREY, 'size': 13})], 15, spacing=6)

# 14 end-of-round questions -----------------------------------------------------------------------------------------
s = slide('每轮结束时问这五个问题', P4)
qs = [('本轮的主要问题解决了吗?', '用数据回答:预测要移走的首个致命错误,实际移走了多少;没移走的原因是什么。'),
      ('哪一步流程又出错了?', '黑板上的纠错记录会被自动收集;同一类错误出现两次以上,就要改流程。'),
      ('哪个错误能做成工具自动拦?', '能写成检查的,就做成工具或门禁,接进流水线,不靠人记。'),
      ('哪些修复可以冻结?', '已在多个应用上验证的公共 API 修复,登记冻结。'),
      ('下一轮先打哪堵墙?', '按最新一轮实测的首个致命错误重新排序,不沿用旧表。')]
for i, (q, a) in enumerate(qs):
    y = 1.4 + i * 1.1
    box(s, 0.5, y, 12.33, 0.95, LIGHT)
    text(s, 0.75, y + 0.22, 4.4, 0.5, q, 17, True, BLUE)
    text(s, 5.3, y + 0.18, 7.4, 0.7, a, 14)

# 15 next -----------------------------------------------------------------------------------------------------------
s = slide('下一轮要验证的')
items = [('U14 收尾', '补测 29 个还没观测过的应用和 3 项欠账,无回退就晋级为新统一态。'),
         ('Wikipedia 文章正文', '在真实应用里用 ArkWeb 渲染 WebView,并验证另外 5 个依赖 WebView 首屏的应用。'),
         ('Flutter 应用', '崩溃点重新定位到 Vulkan 的 Android surface 接口,修法改从 Vulkan 一侧入手。'),
         ('页面大小接口', 'NDK 应用读取页面大小时编号不一致,修复要调整一个冻结项,等待审批;预计 2 到 3 个应用受益。')]
for i, (t, b) in enumerate(items):
    y = 1.4 + i * 1.35
    card(s, 0.5, y, 12.33, 1.15, t, b, PALE['blue'] if i % 2 == 0 else LIGHT, 18, 15)

prs.save(OUT / 'harness-intro.pptx'); print('slides', len(prs.slides))
