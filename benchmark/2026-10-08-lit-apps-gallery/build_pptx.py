#!/usr/bin/env python3
"""Short leader deck (16:9): cover with the counts and 14 representative screenshots, 7 x 2 cards per slide, appendix
tables. Reads report-data.json + img/ written by build_gallery.py from $OUT. Long names use a short display alias
(text before '・' / '：'), approved by the outer loop; the full names are in index.html."""
import json, os, re, sys, unicodedata
from pathlib import Path
sys.path.insert(0, str(Path(__file__).resolve().parents[2] / 'scripts' / 'lab'))
import lab_paths  # noqa: E402
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
from pptx.oxml.ns import qn
OUT = Path(os.environ.get('OUT') or f'{lab_paths.workspaces()}/westlake-inputs/reports/2026-10-08-lit-apps'); d = json.load(open(OUT/'report-data.json'))
FONT = 'PingFang SC'
BLUE, GREY, DARK = RGBColor(0x0a,0x6c,0xff), RGBColor(0x88,0x88,0x88), RGBColor(0x1d,0x1d,0x1f)
CATC = {'效率':'0a6cff','工具':'5856d6','影音':'ff2d55','健康':'34c759','游戏':'ff9500','系统':'8e8e93','图形':'af52de','教育':'00a3a3','通讯':'007aff','阅读':'a2845e'}
CAT_ORDER = list(CATC)
prs = Presentation(); prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
BL = prs.slide_layouts[6]
def text(slide, x, y, w, h, s, size=12, bold=False, color=DARK, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP):
    tb = slide.shapes.add_textbox(x, y, w, h); tf = tb.text_frame; tf.word_wrap = True
    tf.margin_left = tf.margin_right = Emu(0); tf.margin_top = tf.margin_bottom = Emu(0); tf.vertical_anchor = anchor
    for i, line in enumerate(s.split('\n') if isinstance(s, str) else s):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph(); p.alignment = align
        r = p.add_run(); r.text = line; f = r.font; f.size = Pt(size); f.bold = bold; f.color.rgb = color; f.name = FONT
        rPr = r._r.get_or_add_rPr(); ea = rPr.find(qn('a:ea'))
        if ea is None: ea = rPr.makeelement(qn('a:ea'), {}); rPr.append(ea)
        ea.set('typeface', FONT)
    return tb
def header(slide, title, sub=None):
    bar = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, prs.slide_width, Inches(0.08)); bar.fill.solid(); bar.fill.fore_color.rgb = BLUE; bar.line.fill.background()
    text(slide, Inches(0.4), Inches(0.22), Inches(10), Inches(0.5), title, 22, True)
    if sub: text(slide, Inches(0.4), Inches(0.72), Inches(12.5), Inches(0.3), sub, 11, color=GREY)
# cover
s = prs.slides.add_slide(BL); header(s, 'Android 应用在 OpenHarmony 上原包运行 —— 点亮清单',
    'Westlake route A · OpenHarmony 6.1 开发板 · 截图取自历史签认运行记录(未重新拍摄)· 2026-10-08')
kpis = [(str(d['n_pkg']), f"个不同应用的自有界面已签认上屏\n({d['n_keys']} 个测试 APK,其中 Noice 有两个版本)"),
        (str(len(d['shown'])), '本册逐一展示\n首屏内容完整、无报错,每个一张真机截图'),
        ('126 + 3', '测试 APK 中 126 个已进入当前统一态 U13\n3 个(VLC、3D 模型查看器、Termux)为下一版 U14 新增'),
        (f"{len(d['minimal'])} / {len(d['errors'])}", '已点亮但首屏为空列表/加载中(附表)\n/ 首屏带报错框(不计入展示)')]
for i, (n, cap) in enumerate(kpis):
    x, y = Inches(0.4 + (i % 2) * 6.35), Inches(1.35 + (i // 2) * 1.75)
    box = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, x, y, Inches(6.1), Inches(1.5)); box.fill.solid(); box.fill.fore_color.rgb = RGBColor(0xf3,0xf6,0xfb); box.line.fill.background()
    box.adjustments[0] = 0.08
    text(s, x + Inches(0.3), y + Inches(0.15), Inches(5.6), Inches(0.7), n, 36, True, BLUE)
    text(s, x + Inches(0.3), y + Inches(0.85), Inches(5.6), Inches(0.6), cap, 12)
cats = {}
for r in d['shown']: cats[r['cat']] = cats.get(r['cat'], 0) + 1
text(s, Inches(0.4), Inches(4.8), Inches(12.5), Inches(1.2),
     ['“点亮”的定义:未修改的 Android APK 在 OpenHarmony 开发板上安装并启动,启动后第 20 秒截图里是该应用自己的界面,由外环逐张看图签认。',
      '点亮 = 能启动并正常渲染首屏;不代表所有功能都已逐项验证。',
      f"来源(按 {d['n_pkg']} 个应用计):{d['n_pkg']-1} 个来自 F-Droid 开源应用仓库,1 个(VLC)为应用商店原包。卡片上的长应用名用简称(如 Organic Maps),全称见画册 HTML。",
      '展示分类:' + ' · '.join(f'{c} {cats[c]}' for c in CAT_ORDER if c in cats)], 13)
# galleries: 7 x 2 per slide, category order
shown = sorted(d['shown'], key=lambda r: (CAT_ORDER.index(r['cat']), r['name'].lower()))
STRIP = ['wikipedia','fd-organicmaps','vlc','markor','fd-etar','cv2-com-wattwurm-toodoo','cv2-org-andresoviedo-dddmodel2',
         'fd-client','cv2-org-secuso-privacyfriendly2048','anki','fd-tusky','fd-noice','cv2-de-markusfisch-android-shadereditor','fd-uhabits']
byk = {r['key']: r for r in shown}; strip = [byk[k] for k in STRIP if k in byk]
for i, r in enumerate(strip):
    pic = s.shapes.add_picture(str(OUT/'img'/f"{r['key']}.jpg"), Inches(0.4 + i * 0.895), Inches(6.0), height=Inches(1.33))
    pic.line.color.rgb = RGBColor(0xdd,0xdd,0xdd); pic.line.width = Pt(0.5)
PER, COLS = 14, 7
cw, x0, iw = Inches(1.81), Inches(0.42), Inches(1.42)
for pg in range(0, len(shown), PER):
    chunk = shown[pg:pg+PER]; s = prs.slides.add_slide(BL)
    cs = []; [cs.append(r['cat']) for r in chunk if r['cat'] not in cs]
    header(s, '已点亮应用 · ' + ' / '.join(cs), f'第 {pg//PER+1} / {(len(shown)+PER-1)//PER} 页 · 截图为启动后约 20 秒的真机画面')
    for i, r in enumerate(chunk):
        col, row = i % COLS, i // COLS
        x = x0 + cw * col; y = Inches(1.12 + row * 3.18)
        pic = s.shapes.add_picture(str(OUT/'img'/f"{r['key']}.jpg"), x + (cw - iw - Inches(0.1)) // 2, y, width=iw)
        pic.line.color.rgb = RGBColor(0xdd,0xdd,0xdd); pic.line.width = Pt(0.75)
        ty = y + pic.height + Inches(0.05)
        tag = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, x, ty + Inches(0.04), Inches(0.05), Inches(0.2)); tag.fill.solid(); tag.fill.fore_color.rgb = RGBColor.from_string(CATC[r['cat']]); tag.line.fill.background()
        nm = re.split(r'[・：:]', r['name'])[0].strip(); wd = sum(2 if unicodedata.east_asian_width(c) in 'WF' else 1 for c in nm) / 2
        text(s, x + Inches(0.1), ty, cw - Inches(0.2), Inches(0.28), nm, 11 if wd <= 10 else (9.5 if wd <= 11.5 else 8), True)
        text(s, x + Inches(0.1), ty + Inches(0.27), cw - Inches(0.2), Inches(0.42), r['purpose'], 8.5)
# appendix
mins = sorted(d['minimal'], key=lambda r: r['name'].lower())
for part in range(0, len(mins), 16):
    s = prs.slides.add_slide(BL)
    header(s, f"附表 A:已点亮,但首屏没有可展示内容({len(mins)} 个)", '启动正常、显示的是应用自己的界面,但首屏为空列表/加载中/几乎纯色(如空笔记本、空购物清单),截图不适合展示')
    rows = mins[part:part+16]
    tbl = s.shapes.add_table(len(rows)+1, 3, Inches(0.4), Inches(1.15), Inches(12.5), Inches(0.3)*(len(rows)+1)).table
    for j, (h, w) in enumerate([('应用', 2.6), ('是干什么的', 6.2), ('包名', 3.7)]): tbl.columns[j].width = Inches(w)
    for i, row in enumerate([('应用','是干什么的','包名')] + [(r['name'], r['purpose'], r['pkg']) for r in rows]):
        for j, v in enumerate(row):
            c = tbl.cell(i, j); c.text = ''; p = c.text_frame.paragraphs[0]; rr = p.add_run(); rr.text = v
            rr.font.size = Pt(10); rr.font.bold = i == 0; rr.font.name = FONT
            rPr = rr._r.get_or_add_rPr(); ea = rPr.makeelement(qn('a:ea'), {'typeface': FONT}); rPr.append(ea)
            c.margin_top = c.margin_bottom = Emu(20000)
s = prs.slides.add_slide(BL)
header(s, f"附表 B:首屏带报错框({len(d['errors'])} 个,不计入展示)", '应用已启动并显示自己的界面,但首屏弹出错误提示;列出以示完整,不算作展示件')
text(s, Inches(0.4), Inches(1.2), Inches(12.5), Inches(3), [f"{r['name']}({r['pkg']}):{r['purpose']}\n    首屏情况:{r.get('why','')}" for r in d['errors']], 13)
prs.save(OUT/'点亮应用清单.pptx'); print('slides', len(prs.slides))
