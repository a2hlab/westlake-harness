import base64
import collections
from datetime import datetime
import hashlib
import json
from pathlib import Path
import re

HERE = Path(__file__).resolve().parent
WORKSPACES = HERE.parents[2]
MASTER = WORKSPACES / 'westlake-harness'
REPORT = WORKSPACES / 'westlake-inputs/reports/2026-10-08-lit-apps'


def sha(path):
    return hashlib.sha256(path.read_bytes()).hexdigest()


def portable(value):
    if isinstance(value, dict):
        return {portable(key): portable(item) for key, item in value.items()}
    if isinstance(value, list):
        return [portable(item) for item in value]
    if isinstance(value, str):
        return value.replace(str(WORKSPACES) + '/', '$WORKSPACES/').replace(str(Path.home()) + '/', '$HOME/')
    return value


def path(value):
    return Path(re.sub(r'\s+\*gain.*$', '', value).replace('$WORKSPACES/', str(WORKSPACES) + '/').replace('$HOME/', str(Path.home()) + '/')).resolve()


def ownership(row, record, directory_key):
    return [field for field, actual, expected in [
        ('directory', directory_key, row['key']), ('record.key', record['key'], row['key']),
        ('record.package', record['package'], row['pkg'])] if actual != expected]


def coverage(rows, expected):
    counts = collections.Counter(row['key'] for row in rows)
    return dict(total_rows=len(rows), unique_keys=len(counts), missing=sorted(set(expected) - set(counts)),
        extra=sorted(set(counts) - set(expected)), duplicates={key: count for key, count in counts.items() if count > 1})


def metadata_diffs(row, original):
    fields = dict(name='display_name_zh', en='display_name_en', purpose='purpose_zh', pkg='package',
        cat='category', ver='version', source='source')
    return [dict(key=row['key'], field=field, actual=row[field], expected=original[source])
        for field, source in fields.items() if row[field] != original[source]]


def pair_mismatches(actual, expected):
    return dict(missing=list((collections.Counter(expected) - collections.Counter(actual)).elements()),
        extra=list((collections.Counter(actual) - collections.Counter(expected)).elements()))


def signed_membership(key, shot, allowed, original):
    proof = list(allowed.get((key, str(shot)), []))
    if shot == path(original['screenshot_reference']):
        proof.append(dict(metadata_reference=str(MASTER / 'benchmark/2026-10-08-lit-apps-report/apps.json'),
            signoff=original['signoff_path'], scope=original['signed_scope']))
    return proof


def main():
    from lxml import html
    from pptx import Presentation
    from pptx.enum.shapes import MSO_SHAPE_TYPE
    import pptx

    source_paths = [REPORT / name for name in ['report-data.json', 'index.html', '点亮应用清单.pdf', '点亮应用清单.pptx']]
    apps_path = MASTER / 'benchmark/2026-10-08-lit-apps-report/apps.json'
    source_paths.append(apps_path)
    source_before = {str(source): sha(source) for source in source_paths}
    report = json.loads((REPORT / 'report-data.json').read_text())
    apps = {row['key']: row for row in json.loads(apps_path.read_text())['apps']}
    rows = [dict(row, group=group) for group in ['shown', 'minimal', 'errors'] for row in report[group]]
    findings, receipts, differences = [], [], []
    allowed = collections.defaultdict(list)
    signed_sources = []
    for window in range(1, 6):
        directory = MASTER / f'benchmark/2026-10-03-u14-stage-a-plan/window{window}-signoff'
        signoff_path = directory / 'outer-signoff.json'
        sheet_path = directory / 'pair-sheet.json'
        signoff = json.loads(signoff_path.read_text())
        sheet = json.loads(sheet_path.read_text())
        signed_sources += [dict(path=str(source), sha256=sha(source)) for source in [signoff_path, sheet_path]]
        for pair in sheet['pairs']:
            key = pair['key']
            if pair.get('window_t20') and key in signoff['keys']:
                decision = signoff['keys'][key]
                if decision.get('lit_runs', 0) > 0:
                    allowed[(key, str(path(pair['window_t20'])))].append(dict(signoff=str(signoff_path),
                        pair_sheet=str(sheet_path), pair_index=pair['idx'], decision=decision))
    for row in rows:
        original = apps[row['key']]
        differences += metadata_diffs(row, original)
        shot = path(row['shot'])
        record_path = shot.parent / 'record.json'
        record = json.loads(record_path.read_text())
        errors = ownership(row, record, shot.parent.name)
        proof = signed_membership(row['key'], shot, allowed, original)
        if not proof:
            errors.append('unsigned-or-unmapped-run')
        shot_sha = sha(shot)
        matching_shots = [item for item in record.get('screenshots', []) if path(item['path']) == shot]
        if len(matching_shots) != 1 or matching_shots[0].get('sha256') != shot_sha:
            errors.append('record-screenshot-sha-mismatch')
        if record['apk_sha256'] != original['apk_sha256']:
            errors.append('APK-version-sha-mismatch')
        receipts.append(dict(key=row['key'], group=row['group'], shot=str(shot), shot_sha256=shot_sha,
            record=str(record_path), record_sha256=sha(record_path), record_key=record['key'],
            record_package=record['package'], directory_key=shot.parent.name, signed_proof=proof, errors=errors))
        if errors:
            findings.append(dict(id='OBJECT-' + row['key'], severity='high', kind='wrong-object-or-provenance', key=row['key'], reasons=errors))
    partition = coverage(rows, apps)
    if partition['missing'] or partition['extra'] or partition['duplicates']:
        findings.append(dict(id='KEY-COVERAGE', severity='high', kind='requested-key-partition-failed', evidence=partition,
            recommendation='Preserve an explicit alias/coverage row for noice and state package/key units. Do not silently relabel its version as minimal or error just to close counts.'))
    if differences:
        findings.append(dict(id='METADATA', severity='high', kind='metadata-drift', evidence=differences))

    root = html.fromstring((REPORT / 'index.html').read_bytes())
    kpis = [dict(value=node.xpath('string(b)'), label=node.xpath('string(span)')) for node in root.xpath('//div[@class="kpi"]')]
    categories = dict(collections.Counter(row['cat'] for row in report['shown']))
    headers = {node.text.strip(): int(node.xpath('string(span)')) for node in root.xpath('//h2[@class="cat"]')}
    html_cards, html_errors = [], []
    for node in root.xpath('//div[@class="card"]'):
        name_node = node.xpath('.//div[@class="n"]')[0]
        name = (name_node.text or '').strip()
        purpose = node.xpath('string(.//div[@class="p"])')
        metadata = node.xpath('string(.//div[@class="m"])')
        matches = [row for row in report['shown'] if metadata.split(' · ')[0] == row['pkg']]
        if len(matches) != 1:
            html_errors.append(dict(name=name, metadata=metadata, error='ambiguous package/card'))
            continue
        row = matches[0]
        image = node.xpath('.//img/@src')[0]
        embedded = base64.b64decode(image.split(',', 1)[1])
        thumbnail = REPORT / 'img' / (row['key'] + '.jpg')
        error = []
        if (name, purpose, node.get('data-cat')) != (row['name'], row['purpose'], row['cat']):
            error.append('caption/category mismatch')
        html_cards.append(dict(key=row['key'], name=name, purpose=purpose, category=node.get('data-cat'),
            thumbnail_sha256=sha(thumbnail), embedded_sha256=hashlib.sha256(embedded).hexdigest(),
            thumbnail_bytes_identical=embedded == thumbnail.read_bytes(), errors=error))
        if error:
            html_errors.append(dict(key=row['key'], errors=error))
    table_counts = [len(table.xpath('.//tr')) - 1 for table in root.xpath('//table')]
    totals = dict(packages=len({row['pkg'] for row in rows}), shown=len(report['shown']), minimal=len(report['minimal']),
        errors=len(report['errors']), reference=126, gains=3, requested_keys=len(apps))
    cover_errors = []
    if [item['value'] for item in kpis] != ['128', '94', '126 + 3', '32 / 2']:
        cover_errors.append('unexpected KPI values')
    if report['n_keys'] != len(apps) or report['n_pkg'] != totals['packages']:
        cover_errors.append('declared key/package totals mismatch')
    if '129 个测试 APK' not in kpis[0]['label']:
        cover_errors.append('129-key cover scope missing')
    if headers != categories or sum(headers.values()) != len(report['shown']):
        cover_errors.append('category counts not closed over shown')
    if table_counts != [len(report['minimal']), len(report['errors'])]:
        cover_errors.append('appendix row count mismatch')
    if len(html_cards) != len(report['shown']) or html_errors:
        cover_errors.append('HTML card content mismatch')
    expected_keys_by_html = [row['key'] for row in html_cards]
    if collections.Counter(expected_keys_by_html) != collections.Counter(row['key'] for row in report['shown']):
        cover_errors.append('HTML shown key set differs')
    definition = root.xpath('string(//div[@class="def"])')
    fdroid_included = sum('F-Droid' in row['source'] for row in rows)
    if '来源:128 个来自 F-Droid' in definition and fdroid_included != 128:
        findings.append(dict(id='SOURCE-COUNT-UNITS', severity='medium', kind='cover-source-unit-ambiguity',
            evidence=dict(text='来源:128 个来自 F-Droid 开源仓库,1 个(VLC)为应用商店原包',
                included_packages_fdroid=fdroid_included, included_packages_other=len(rows) - fdroid_included,
                original_key_scope_fdroid=128, original_key_scope_other=1),
            recommendation='Label 128+1 explicitly as original APK/key source counts, or use 127+1 for the 128 represented packages.'))
    if cover_errors:
        findings.append(dict(id='HTML', severity='high', kind='html-closure-failed', evidence=cover_errors + html_errors))

    presentation = Presentation(str(REPORT / '点亮应用清单.pptx'))
    ppt_cards, ppt_pages, ppt_pair_errors = [], [], []
    thumbnail_keys = collections.defaultdict(list)
    for row in report['shown']:
        thumbnail_keys[sha(REPORT / 'img' / (row['key'] + '.jpg'))].append(row['key'])
    for number, slide in enumerate(presentation.slides, 1):
        texts = [shape for shape in slide.shapes if shape.has_text_frame and shape.text]
        pictures = [shape for shape in slide.shapes if shape.shape_type == MSO_SHAPE_TYPE.PICTURE]
        page_cards = []
        for picture in sorted(pictures, key=lambda shape: (shape.top, shape.left)):
            if number == 1:
                continue
            candidates = [shape for shape in texts if picture.left - 100000 <= shape.left <= picture.left + picture.width
                and 0 <= shape.top - (picture.top + picture.height) < 900000]
            candidates.sort(key=lambda shape: shape.top)
            if len(candidates) != 2 or candidates[0].left != candidates[1].left:
                ppt_pair_errors.append(dict(page=number, error='ambiguous card text geometry', texts=[shape.text for shape in candidates]))
                continue
            name, purpose = [shape.text for shape in candidates]
            image_sha = hashlib.sha256(picture.image.blob).hexdigest()
            matches = thumbnail_keys[image_sha]
            key = matches[0] if len(matches) == 1 else None
            card = dict(page=number, name=name, purpose=purpose, key=key, x=picture.left, y=picture.top,
                image_sha256=image_sha)
            if key is None:
                ppt_pair_errors.append(dict(page=number, name=name, purpose=purpose, error='image does not identify one shown key'))
            else:
                expected = next(row for row in report['shown'] if row['key'] == key)
                if (name, purpose) != (expected['name'], expected['purpose']):
                    ppt_pair_errors.append(dict(page=number, key=key, name=name, purpose=purpose,
                        expected_name=expected['name'], expected_purpose=expected['purpose'], error='caption differs from shown'))
            page_cards.append(card)
        ppt_cards += page_cards
        ppt_pages.append(dict(page=number, pictures=len(pictures), cards=page_cards))
    pair_diff = pair_mismatches([(row['name'], row['purpose']) for row in ppt_cards],
        [(row['name'], row['purpose']) for row in report['shown']])
    if pair_diff['missing'] or pair_diff['extra'] or ppt_pair_errors:
        findings.append(dict(id='PPT-CARDS', severity='high', kind='ppt-card-pair-mismatch', evidence=pair_diff, errors=ppt_pair_errors))

    source_after = {str(source): sha(source) for source in source_paths}
    if source_after != source_before:
        findings.append(dict(id='SOURCE-DRIFT', severity='high', kind='outer-files-changed-during-audit'))
    checks = dict(shot_object_and_signed_run='pass' if all(not row['errors'] for row in receipts) else 'fail',
        metadata_exact='pass' if not differences else 'fail',
        union129='pass' if not any(partition[field] for field in ['missing', 'extra', 'duplicates']) else 'fail',
        shown_noice_once='pass' if sum(row['pkg'] == 'com.github.ashutoshgngwr.noice' for row in report['shown']) == 1 else 'fail',
        html_package_counts_categories_and_appendices='pass' if not cover_errors else 'fail',
        ppt_card_name_purpose_pairs='pass' if not (pair_diff['missing'] or pair_diff['extra'] or ppt_pair_errors) else 'fail')
    result = dict(schema='lit-gallery-mechanical-check/v1', generated_at=datetime.now().astimezone().isoformat(),
        visual_judgment=False, board_access=False, outer_files_modified=False,
        metadata_commit='c89296a4c6', source_sha256=source_before, sources_unchanged=source_after == source_before,
        checks=checks, findings=findings, totals=totals, partition=partition, field_differences=differences,
        screenshot_receipts=receipts, signed_sources=signed_sources, html=dict(kpis=kpis, categories=headers, appendix_rows=table_counts,
            cards=html_cards, errors=html_errors), ppt=dict(reader='python-pptx', version=pptx.__version__,
            slide_count=len(presentation.slides), card_count=len(ppt_cards), pages=ppt_pages, pair_differences=pair_diff,
            errors=ppt_pair_errors), limits=['No pixel or rendering judgment', 'PDF bound by SHA only; no PDF content audit requested',
                'Cover current unified-state/complete-first-screen assertions not independently revalidated',
                'PPT thumbnails matched by exact bytes; HTML uses different image bytes, recorded without judging image derivation or pixels',
                'PPT cover collage is not a gallery card and excluded from card checks'])
    HERE.mkdir(parents=True, exist_ok=True)
    (HERE / 'findings.json').write_text(json.dumps(portable(result), ensure_ascii=False, indent=2) + '\n')
    print(json.dumps(dict(checks=checks, finding_ids=[finding['id'] for finding in findings], totals=totals,
        ppt_cards=len(ppt_cards), sources_unchanged=result['sources_unchanged']), ensure_ascii=False))


if __name__ == '__main__':
    main()
