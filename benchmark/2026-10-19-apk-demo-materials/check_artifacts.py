"""Inspect real delivery packages and fail on corruption, missing evidence or false counts."""
import base64,csv,hashlib,io,json,posixpath,re,sys,zipfile,subprocess
from collections import Counter
from pathlib import Path
from xml.etree import ElementTree as ET
from PIL import Image
from pptx import Presentation
from openpyxl import load_workbook
H=Path(__file__).resolve().parent;O=H/'delivery'
def sha(p):return hashlib.sha256(p.read_bytes()).hexdigest()
def check_zip(path):
 with zipfile.ZipFile(path) as z:
  assert z.testzip() is None,f'Corrupt archive: {path}'
  names=set(z.namelist())
  for name in names:
   assert not name.startswith('/') and '..' not in Path(name).parts
   if not name.endswith('.rels'):continue
   base=posixpath.dirname(posixpath.dirname(name)) if name!='_rels/.rels' else ''
   for rel in ET.fromstring(z.read(name)):
    if rel.get('TargetMode')=='External':continue
    target=posixpath.normpath(posixpath.join(base,rel.get('Target'))).lstrip('/')
    assert target in names,(path,name,target)
 return names

def main():
 d=json.loads((O/'manifest.json').read_text());rows=d['rows'];by={r['key']:r for r in rows}
 assert len(rows)==len(by)==130
 counts=Counter(r['status'] for r in rows);assert counts=={'startup-only':115,'excluded':13,'scoped-special':2},counts
 assert len([r for r in rows if r['scope']=='baseline128'])==128
 source=json.loads((H/'evidence/source-demo-list.json').read_text())
 assert {r['key'] for r in rows if r['status']=='startup-only'}==set(source['ordinary_app_keys'])
 assert {r['key'] for r in rows if r['status']=='excluded'}=={r['key'] for r in source['excluded_keys_with_reasons']}
 for r in rows:
  assert r['name'] and r['purpose'] and r['source_image'] and r['source_record']
  p=O/r['image'];assert sha(p)==r['image_sha256'],r['key']
  with Image.open(p) as im:assert im.size==(r['width'],r['height']);im.verify()
  assert r['status']!='excluded' or r['exclusion_reason']
 selected=json.loads((O/'selected.json').read_text())['keys'];assert len(selected)==16 and selected[:2]==['toutiao-route-b','boatattack']
 ppt_reports=[]
 for name,expected in [('Westlake-精选演示-17页.pptx',selected),('Westlake-全量分类-128加2.pptx',list(by))]:
  p=O/name;check_zip(p);prs=Presentation(p);assert abs(prs.slide_width/prs.slide_height-16/9)<.0001
  found=[]
  for s in prs.slides:
   for sh in s.shapes:
    assert sh.left>=0 and sh.top>=0 and sh.left+sh.width<=prs.slide_width+5 and sh.top+sh.height<=prs.slide_height+5,(name,sh.name,'off canvas')
    if sh.shape_type==13:
     key=sh.name.removeprefix('evidence:');assert key in by,key;r=by[key]
     assert hashlib.sha256(sh.image.blob).hexdigest()==r['image_sha256'],key
     assert abs((sh.width/sh.height)/(r['width']/r['height'])-1)<.00001,(key,'stretch')
     assert sh.crop_left==sh.crop_right==sh.crop_top==sh.crop_bottom==0,(key,'crop')
     found.append(key)
  assert Counter(found)==Counter(expected),(name,'key coverage')
  if '精选' in name:assert len(prs.slides)==17
  ppt_reports.append({'path':name,'slides':len(prs.slides),'pictures':len(found),'sha256':sha(p)})
 with (O/'全量清单-128加2.csv').open(encoding='utf-8-sig',newline='') as f:csvrows=list(csv.DictReader(f))
 assert {r['key'] for r in csvrows}==set(by) and len(csvrows)==130
 x=O/'全量图文清单-128加2.xlsx';check_zip(x);wb=load_workbook(x);ws=wb['全部130项'];assert ws.max_row==131 and len(ws._images)==130
 assert {ws.cell(i,10).value for i in range(2,132)}==set(by)
 html=(O/'全量图文清单.html').read_text();assert html.count('<article ')==130 and html.count('data:image/jpeg;base64,')==130
 for ref in re.findall(r'href="(originals/[^\"]+)"',html):assert (O/ref).exists()
 raw=json.loads((O/'raw-baseline-index.json').read_text())['screenshots']
 assert len(raw)==256 and Counter(r['key'] for r in raw)=={r['key']:2 for r in rows if r['scope']=='baseline128'}
 for r in raw:assert r['captured'] is True and sha(O/r['image'])==r['sha256']
 for stem,count in [('Westlake-精选演示-17页',17),('Westlake-全量分类-128加2',50)]:
  pdf=O/(stem+'.pdf');assert pdf.is_file(),pdf
  info=subprocess.check_output(['pdfinfo',str(pdf)],text=True);assert re.search(r'Pages:\s+'+str(count)+r'\b',info),info
  if count==17:
   text=subprocess.check_output(['pdftotext',str(pdf),'-'],text=True)
   assert '记录桌面游戏中的玩家生命值与分数' in text.replace('\n',''), 'OneTwo PDF stale'
 assert len([p for p in O.glob('*.pptx') if '兼容版' not in p.name])==2 and len(list(O.glob('*.pdf')))==2, 'Two canonical deck pairs; preserve previously shared alias'
 full=O/'Westlake-完整离线包.zip'
 if full.exists():
  check_zip(full)
  with zipfile.ZipFile(full) as z:
   assert len([n for n in z.namelist() if n.endswith('.pptx')])==2 and len([n for n in z.namelist() if n.endswith('.pdf')])==2
   assert not any('兼容版' in n for n in z.namelist())
   for r in rows:assert hashlib.sha256(z.read(r['image'])).hexdigest()==r['image_sha256']
   for r in raw:assert hashlib.sha256(z.read(r['image'])).hexdigest()==r['sha256']
   for p in [O/a['path'] for a in ppt_reports]+[x,O/'manifest.json',O/'raw-baseline-index.json',O/'全量图文清单.html',O/'全量清单-128加2.csv']+list(O.glob('*.pdf')):
    assert hashlib.sha256(z.read(p.name)).hexdigest()==sha(p),('stale bundle',p.name)
 report={'status':'PASS_INTEGRITY','baseline_keys':128,'rows':130,'status_counts':dict(counts),'original_images_verified':130,'raw_baseline_captured_images_verified':256,'xlsx_embedded_images':130,'pptx':ppt_reports,'render_review':'separate visual review required','facts_total_verbatim':(H/'evidence/facts.txt').read_text().strip().splitlines()[-1]}
 (H/'results.json').write_text(json.dumps(report,ensure_ascii=False,indent=2));print(json.dumps(report,ensure_ascii=False))
if __name__=='__main__':main()
