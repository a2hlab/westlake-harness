"""Generate portable illustrated inventory from the reviewed manifest, no device access."""
import base64, csv, html, io, json, zipfile
from pathlib import Path
from PIL import Image
from openpyxl import Workbook
from openpyxl.drawing.image import Image as XLImage
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter
H=Path(__file__).resolve().parent;O=H/'delivery';D=json.loads((O/'manifest.json').read_text());rows=D['rows'];thumb=O/'thumbnails';thumb.mkdir(exist_ok=True)
STATUS={'startup-only':'仅启动亮屏','excluded':'排除，不计入演示','scoped-special':'专项范围内验证'}
fields=['key','name','purpose','category','status','verified_scope','exclusion_reason','observed_limitations','image','image_sha256','source_image','source_record','date','scope','package']
with (O/'全量清单-128加2.csv').open('w',encoding='utf-8-sig',newline='') as f:
 w=csv.DictWriter(f,fieldnames=fields,extrasaction='ignore');w.writeheader();w.writerows(rows)
wb=Workbook();readme=wb.active;readme.title='阅读说明'
for line in ['Westlake 全量图文清单','128 个基准键 = 115 startup-only + 13 excluded；另加 2 项专项，共 130 行。','亮屏不等于完整功能可用。用途说明是产品用途，verified_scope 是实测范围。','头条采用 2026-10-19 信息流原图；BoatAttack 采用 2026-10-14 复位后原图。','全部截图为原始文件的等比例缩略图；原图在 originals，路径与 SHA256 可追溯。','排除行保留失败、加载或闪屏画面作为证据，不计入成功演示。','原始运行事实：'+(H/'evidence/facts.txt').read_text().strip().splitlines()[-1]]:readme.append([line])
readme.column_dimensions['A'].width=135
for row in readme:
 row[0].alignment=Alignment(wrap_text=True,vertical='center');readme.row_dimensions[row[0].row].height=38
ws=wb.create_sheet('全部130项')
headers=['序号','真实截图','应用名称','主要用途','分类','验证状态','实测范围','排除原因','观察限制（原文）','基准键','包名','截图日期','原图相对路径','SHA256','原图来源','记录来源','分组']
ws.append(headers);widths=[7,30,28,40,13,23,40,45,60,50,45,16,55,68,90,90,20]
for i,(head,width) in enumerate(zip(headers,widths),1):
 ws.column_dimensions[get_column_letter(i)].width=width;cell=ws.cell(1,i);cell.fill=PatternFill('solid',fgColor='101B2B');cell.font=Font(color='FFFFFF',bold=True);cell.alignment=Alignment(wrap_text=True)
ws.row_dimensions[1].height=30;ws.freeze_panes='D2';ws.auto_filter.ref=f'A1:Q{len(rows)+1}'
cards=[]
for i,r in enumerate(rows,2):
 im=Image.open(O/r['image']);im.thumbnail((360,260));tp=thumb/(r['key']+'.jpg');im.convert('RGB').save(tp,quality=86)
 ws.append([i-1,'',r['name'],r['purpose'],r['category'],STATUS[r['status']],r['verified_scope'],r['exclusion_reason'],r['observed_limitations'],r['key'],r['package'],r['date'],r['image'],r['image_sha256'],r['source_image'],r['source_record'],r['scope']])
 for cell in ws[i]:cell.alignment=Alignment(wrap_text=True,vertical='top');cell.font=Font(name='PingFang SC',size=11)
 ws.row_dimensions[i].height=156
 img=XLImage(tp);scale=min(190/img.width,195/img.height);img.width*=scale;img.height*=scale;ws.add_image(img,f'B{i}')
 ws.cell(i,13).hyperlink=r['image'];ws.cell(i,13).style='Hyperlink'
 if r['status']=='excluded':
  for cell in ws[i]:cell.fill=PatternFill('solid',fgColor='FFF0E5')
 b64=base64.b64encode(tp.read_bytes()).decode();e=html.escape
 cards.append(f'''<article data-cat="{e(r['category'])}" data-status="{r['status']}" data-query="{e(' '.join([r['key'],r['name'],r['purpose'],r['package']]).lower())}"><a href="{e(r['image'])}" target="_blank"><img src="data:image/jpeg;base64,{b64}" alt="{e(r['name'])}真实截图"></a><div><span class="badge {r['status']}">{e(STATUS[r['status']])}</span><h2>{e(r['name'])}</h2><p>{e(r['purpose'])}</p><p class="scope">实测范围：{e(r['verified_scope'])}</p><p>{e(r['exclusion_reason'] or '')}</p><details><summary>证据与观察限制</summary><p>{e(r['observed_limitations'])}</p><p>{e(r['key'])} / {e(r['package'])}</p><p>截图日期：{r['date']}</p><p>原图来源：{e(r['source_image'])}</p><p>SHA256：{r['image_sha256']}</p></details></div></article>''')
wb.save(O/'全量图文清单-128加2.xlsx')
opts=''.join(f'<option>{html.escape(c)}</option>' for c in sorted({r['category'] for r in rows}))
doc='''<!DOCTYPE html><html lang="zh-CN"><meta charset="utf-8"><meta name="viewport" content="width=device-width, initial-scale=1"><title>Westlake 全量图文清单</title><style>
*{box-sizing:border-box}body{margin:0;font:16px -apple-system,"PingFang SC",sans-serif;background:#eef2f6;color:#152238}header{background:#101b2b;color:white;padding:40px 5vw}h1{font-size:34px}header p{max-width:1050px;line-height:1.7;color:#c7d5e5}.toolbar{position:sticky;top:0;z-index:1;display:flex;gap:12px;flex-wrap:wrap;background:#fff;padding:18px 5vw;box-shadow:0 1px 8px #0001}input,select{font:inherit;border:1px solid #cad2dd;padding:10px;border-radius:6px}main{padding:25px 5vw;display:grid;grid-template-columns:repeat(auto-fit,minmax(340px,1fr));gap:22px}article{background:white;border-radius:10px;overflow:hidden;border:1px solid #dce3eb}article>a{display:block;background:#e2e7ed;text-align:center;padding:18px}img{width:100%;height:300px;object-fit:contain}article>div{padding:22px}h2{font-size:24px;margin:14px 0 8px}p{line-height:1.55}details{font-size:12px;overflow-wrap:anywhere;color:#526276}.badge{font-size:12px;background:#e7edf5;padding:5px 9px;border-radius:12px}.excluded{background:#ffe5d2;color:#893607}.scoped-special{background:#d8f3eb;color:#0a6853}.scope{color:#405773;font-size:14px}footer{padding:30px 5vw;color:#526276}</style>
<header><h1>Android 应用在 OpenHarmony 上运行</h1><p>全量图文清单 · 128 个基准键 + 2 项专项</p><p><b>115 项仅启动亮屏 · 13 项排除 · 2 项专项另计。</b>用途是应用的产品用途，实测范围另列。欢迎、空状态与错误界面均如实保留；亮屏不代表完整功能可用。点击截图可打开未改动原图。</p></header><div class="toolbar"><input id="search" placeholder="搜索名称 / 用途 / key"><select id="cat"><option value="">全部分类</option>'''+opts+'''</select><select id="status"><option value="">全部状态</option><option value="startup-only">仅启动亮屏</option><option value="excluded">排除项</option><option value="scoped-special">专项证据</option></select><span id="count"></span></div><main>'''+''.join(cards)+'''</main><footer>历史证据整理于 2026-10-19。全程未操作设备。离线文件：manifest.json、CSV、XLSX、PPTX 和 originals/。</footer><script>const q=document.querySelector('#search'),c=document.querySelector('#cat'),s=document.querySelector('#status');function filter(){let n=0;document.querySelectorAll('article').forEach(a=>{const ok=(!q.value||a.dataset.query.includes(q.value.toLowerCase()))&&(!c.value||a.dataset.cat===c.value)&&(!s.value||a.dataset.status===s.value);a.hidden=!ok;if(ok)n++});document.querySelector('#count').textContent=n+' / 130 项';}q.oninput=c.onchange=s.onchange=filter;filter();</script></html>'''
(O/'全量图文清单.html').write_text(doc)
print('130-row XLSX with 130 embedded screenshots, CSV and offline HTML generated')
